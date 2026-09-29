"""Deterministic ingestion and evidence budgeting; safe to test without infrastructure."""
from __future__ import annotations
import io,re
from pathlib import Path
from docx import Document as DocxDocument
from pypdf import PdfReader

def estimate_tokens(text:str)->int:
    # Cheap proxy (~0.75 words/token); benchmark reports this as an estimate.
    return max(1,int(len(re.findall(r'\w+|[^\w\s]',text))/0.75))

class DocumentExtractor:
    def extract(self,data:bytes,filename:str)->list[tuple[int,str]]:
        suffix=Path(filename).suffix.lower()
        if suffix=='.pdf':return [(i+1,(p.extract_text() or '').strip()) for i,p in enumerate(PdfReader(io.BytesIO(data)).pages)]
        if suffix=='.docx':
            d=DocxDocument(io.BytesIO(data)); blocks=[p.text.strip() for p in d.paragraphs if p.text.strip()]
            blocks += [' | '.join(c.text.strip() for c in row.cells) for table in d.tables for row in table.rows]
            return [(1,'\n'.join(blocks))]
        if suffix in {'.txt','.md','.html'}:
            text=data.decode('utf-8',errors='replace')
            if suffix=='.html':text=re.sub(r'<[^>]+>',' ',text)
            return [(1,text)]
        raise ValueError('Supported formats: PDF, DOCX, TXT, Markdown and HTML. PPTX is not enabled.')

class StructureChunker:
    """Sentence and paragraph aware bounded chunks with controlled overlap."""
    def __init__(self,max_words=420,overlap=65):self.max_words=max_words;self.overlap=overlap
    def split(self,text:str,page:int):
        paras=[re.sub(r'\s+',' ',p).strip() for p in re.split(r'\n{1,}|(?<=[.!?])\s+(?=[A-Z0-9])',text) if p.strip()]
        out=[];buf=[];n=0
        def flush():
            nonlocal buf,n
            if buf:out.append((page,' '.join(buf)))
            buf=' '.join(buf).split()[-self.overlap:];n=len(buf)
        for para in paras:
            words=para.split()
            if len(words)>self.max_words:
                if buf:flush()
                for i in range(0,len(words),self.max_words-self.overlap):out.append((page,' '.join(words[i:i+self.max_words])))
                buf=[];n=0;continue
            if n+len(words)>self.max_words:flush()
            buf.extend(words);n+=len(words)
        if buf:out.append((page,' '.join(buf)))
        return [x for x in out if len(x[1])>20]

def pack_ranked_passages(passages,budget:int,max_passages:int):
    """Pack ranked tuples(score, chunk, document) within a token proxy budget."""
    packed=[];used=0;seen=set();naive=max(1,sum(p[1].token_estimate for p in passages[:max_passages]))
    for score,chunk,doc in passages:
        fingerprint=re.sub(r'\W+',' ',chunk.text.lower())[:600]
        if fingerprint in seen:continue
        seen.add(fingerprint)
        remain=budget-used
        if remain<=0 or len(packed)>=max_passages:break
        content=chunk.text
        if chunk.token_estimate>remain:content=content[:max(120,int(remain*3.4))].rsplit(' ',1)[0]
        tokens=estimate_tokens(content)
        if tokens<20:continue
        packed.append({'rank':len(packed)+1,'score':round(float(score),4),'document_id':doc.id,'title':doc.title,'version':doc.version,'page':chunk.page,'chunk_id':chunk.id,'text':content,'snapshot_url':f'/documents/{doc.id}/snapshot?page={chunk.page}'})
        used+=tokens
    return packed,used,naive,max(0,round(100*(1-used/naive),1))
